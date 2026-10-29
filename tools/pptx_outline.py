#!/usr/bin/env python3
"""Строит индекс слайдов из .pptx для docs/content/slides.md.

Зачем: содержание презентации должно быть читаемым текстом в репозитории — иначе
единственный способ узнать, что внутри бинарного файла, — открыть PowerPoint.

    python tools/pptx_outline.py "Концепция_счастья.pptx" > docs/content/slides.md
"""
from __future__ import annotations

import sys
from pathlib import Path


def iter_text(shape, acc: list[str]) -> None:
    """Собирает текст, включая вложенные группы фигур."""
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for child in shape.shapes:
            iter_text(child, acc)
        return
    if shape.has_text_frame:
        text = shape.text.strip()
        if text:
            acc.append(" ".join(text.replace("\v", " ").split()))


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    from pptx import Presentation

    source = Path(argv[1])
    deck = Presentation(source)
    print(f"# Содержание презентации: {source.name}\n")
    print(f"Слайдов: **{len(deck.slides)}**. Индекс построен автоматически из файла:")
    print("`python tools/pptx_outline.py` — не редактируйте этот файл вручную.\n")
    for number, slide in enumerate(deck.slides, 1):
        shapes: list[str] = []
        for shape in slide.shapes:
            iter_text(shape, shapes)
        title = shapes[0] if shapes else "(без текста)"
        details = [part for part in shapes[1:] if len(part) < 160][:6]
        print(f"## {number}. {title}\n")
        if details:
            print("Тезисы на слайде:\n")
            for part in details:
                print(f"- {part}")
            print()
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
            if notes:
                print(f"<details><summary>Текст доклада (заметки)</summary>\n\n{notes}\n\n</details>\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
